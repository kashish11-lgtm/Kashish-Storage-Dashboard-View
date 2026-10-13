"""Add a 'Lubes Deep Dive' tab to the Talabat Durex workbook.

Numbers are formulas pointing at 'Lubes SKU level summary' (the user's pivot),
except store-reach counts, which come from 'Raw data' Location (only populated
for Jan/Apr/May/Aug 2025) and are hardcoded with a source note.
"""
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as CL

src, out = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(src)
NAME = 'Lubes Deep Dive'
if NAME in wb.sheetnames:
    del wb[NAME]
ws = wb.create_sheet(NAME, index=wb.sheetnames.index('Lubes SKU level summary') + 1)
S = "'Lubes SKU level summary'!"

FONT = 'Aptos Narrow'
f_norm = Font(name=FONT, size=11)
f_bold = Font(name=FONT, size=11, bold=True)
f_title = Font(name=FONT, size=14, bold=True)
f_sec = Font(name=FONT, size=12, bold=True, color='FFFFFF')
f_input = Font(name=FONT, size=11, color='0000FF')
f_note = Font(name=FONT, size=9, italic=True, color='595959')
fill_sec = PatternFill('solid', fgColor='1F3864')
fill_hdr = PatternFill('solid', fgColor='D9E1F2')
fill_tot = PatternFill('solid', fgColor='F2F2F2')
fill_in = PatternFill('solid', fgColor='FFFF00')
fill_red = PatternFill('solid', fgColor='F8CBAD')
fill_grn = PatternFill('solid', fgColor='C6EFCE')
thin = Side(style='thin', color='BFBFBF')
box = Border(top=thin, bottom=thin, left=thin, right=thin)
wrap = Alignment(wrap_text=True, vertical='top')
center = Alignment(horizontal='center', vertical='center', wrap_text=True)

AED = '#,##0;(#,##0);-'
PCT = '0.0%;(0.0%);-'
NUM = '#,##0;(#,##0);-'
DEC = '#,##0.0;(#,##0.0);-'


def put(cell, v, font=f_norm, fmt=None, fill=None, align=None, border=True):
    c = ws[cell]
    c.value = v
    c.font = font
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if align:
        c.alignment = align
    if border:
        c.border = box
    return c


def section(row, text, width=16):
    for col in range(2, 2 + width):
        ws.cell(row=row, column=col).fill = fill_sec
    c = ws.cell(row=row, column=2, value=text)
    c.font = f_sec


def header(row, labels, start_col=2):
    for i, lab in enumerate(labels):
        put(f'{CL(start_col + i)}{row}', lab, f_bold, fill=fill_hdr, align=center)
    ws.row_dimensions[row].height = 45


ws['B1'] = 'Durex Lubes on Talabat: SKU deep dive'
ws['B1'].font = f_title
ws['B2'] = ("All numbers are formulas on 'Lubes SKU level summary' unless marked. Revenue is in the "
            "same unit as the source data. JA = Jan-Apr. H2'25 = Jul-Dec 2025 monthly average.")
ws['B2'].font = f_note

# ---------------------------------------------------------------- A. scorecard
r = 4
section(r, 'A. SKU scorecard: YoY growth vs. the H2-2025 run rate')
header(r + 1, ['SKU', 'Pack', "Rev JA'25", "Rev JA'26", 'YoY rev %', "Units JA'25", "Units JA'26",
               'YoY units %', "ASP JA'25", "ASP JA'26", 'ASP chg %', "H2'25 avg rev / mo",
               "Apr'26 rev", "Apr'26 vs H2'25 avg", "Share of lubes rev JA'26", 'Read'])
skus = [  # (source row, short name, pack, read)
    (8, 'Play Cheeky Cherry', '50 ml', 'Hero SKU, about 1/3 of lubes. Growing and at its H2 level by April.'),
    (9, 'Play Feel', '50 ml', 'Fastest grower (+56% units YoY). Holding its H2 level.'),
    (10, 'Play Tingle', '50 ml', 'Growing YoY but softer than the other Play SKUs. ~10% below its H2 level.'),
    (11, 'Maxima Lubricant Gel', '50 ml', 'The only 50 ml SKU declining. ASP up 14% with no promo since Jan-26 (see B).'),
    (13, 'Play Massage 2-in-1 Aloe Vera', '200 ml', 'Out-sells Ylang but has the biggest fall vs H2 of any SKU.'),
    (14, 'Play Massage 2-in-1 Ylang Ylang', '200 ml', 'YoY flattered by a weak Jan-25 base (only 14 stores ranged it).'),
    (16, 'Naturals Pure Intimate', '100 ml', '17 units in 16 months, 1 store at a time. Effectively not ranged.'),
]
row = r + 2
first = row
for src_row, name, pack, read in skus:
    s = src_row
    put(f'B{row}', name)
    put(f'C{row}', pack, align=Alignment(horizontal='center'))
    put(f'D{row}', f'=SUM({S}C{s}:F{s})', fmt=AED)
    put(f'E{row}', f'=SUM({S}P{s}:S{s})', fmt=AED)
    put(f'F{row}', f'=IFERROR(E{row}/D{row}-1,0)', fmt=PCT)
    put(f'G{row}', f'=SUM({S}U{s}:X{s})', fmt=NUM)
    put(f'H{row}', f'=SUM({S}AH{s}:AK{s})', fmt=NUM)
    put(f'I{row}', f'=IFERROR(H{row}/G{row}-1,0)', fmt=PCT)
    put(f'J{row}', f'=IFERROR(D{row}/G{row},0)', fmt=AED)
    put(f'K{row}', f'=IFERROR(E{row}/H{row},0)', fmt=AED)
    put(f'L{row}', f'=IFERROR(K{row}/J{row}-1,0)', fmt=PCT)
    put(f'M{row}', f'=SUM({S}I{s}:N{s})/6', fmt=AED)
    put(f'N{row}', f'=N({S}S{s})', fmt=AED)
    put(f'O{row}', f'=IFERROR(N{row}/M{row}-1,0)', fmt=PCT)
    put(f'P{row}', f'=IFERROR(E{row}/E${first + len(skus)},0)', fmt=PCT)
    put(f'Q{row}', read, align=wrap)
    ws.row_dimensions[row].height = 30
    row += 1
tot = row
put(f'B{tot}', 'Total lubes', f_bold, fill=fill_tot)
put(f'C{tot}', '', fill=fill_tot)
for col, fmt in zip('DEFGHIJKLMNOP', [AED, AED, PCT, NUM, NUM, PCT, AED, AED, PCT, AED, AED, PCT, PCT]):
    if col in 'DEGHMN':
        v = f'=SUM({col}{first}:{col}{tot - 1})'
    else:
        v = ws[f'{col}{first}'].value.replace(str(first), str(tot)) if col != 'P' else f'=IFERROR(E{tot}/E{tot},0)'
    put(f'{col}{tot}', v, f_bold, fmt=fmt, fill=fill_tot)
put(f'Q{tot}', 'Lubes are growing YoY; the gap is against the H2-25 run rate, not last year.',
    f_bold, fill=fill_tot, align=wrap)
ws.row_dimensions[tot].height = 30
put(f'B{tot + 1}', f"=\"Check vs pivot Grand Total JA'26: \"&TEXT({S}T17-E{tot},\"#,##0\")&\" difference\"",
    f_note, border=False)
ws.conditional_formatting.add(f'F{first}:F{tot}', openpyxl.formatting.rule.CellIsRule(operator='lessThan', formula=['0'], fill=fill_red))
ws.conditional_formatting.add(f'I{first}:I{tot}', openpyxl.formatting.rule.CellIsRule(operator='lessThan', formula=['0'], fill=fill_red))
ws.conditional_formatting.add(f'O{first}:O{tot}', openpyxl.formatting.rule.CellIsRule(operator='lessThan', formula=['-0.1'], fill=fill_red))
ws.conditional_formatting.add(f'L{first}:L{tot}', openpyxl.formatting.rule.CellIsRule(operator='greaterThan', formula=['0.05'], fill=fill_red))

# ---------------------------------------------------------------- B. Maxima price
r = tot + 3
section(r, 'B. Maxima 50 ml: price premium vs Play 50 ml, and its share of 50 ml units', 18)
months = ['Jan-25', 'Feb-25', 'Mar-25', 'Apr-25', 'May-25', 'Jun-25', 'Jul-25', 'Aug-25', 'Sep-25', 'Oct-25',
          'Nov-25', 'Dec-25', 'Jan-26', 'Feb-26', 'Mar-26', 'Apr-26']
rev_cols = [CL(i) for i in range(3, 15)] + ['P', 'Q', 'R', 'S']          # C..N, P..S
unit_cols = [CL(i) for i in range(21, 33)] + ['AH', 'AI', 'AJ', 'AK']   # U..AF, AH..AK
header(r + 1, ['Metric'] + months)
labels = ['Maxima ASP', 'Play 50 ml blended ASP', 'Maxima premium vs Play', 'Maxima share of 50 ml units',
          'Play 50 ml units', 'Maxima units']
for i, lab in enumerate(labels):
    put(f'B{r + 2 + i}', lab, f_bold)
for j, (rc, uc) in enumerate(zip(rev_cols, unit_cols)):
    col = CL(3 + j)
    put(f'{col}{r + 2}', f'=IFERROR({S}{rc}11/{S}{uc}11,0)', fmt=AED)
    put(f'{col}{r + 3}', f'=IFERROR(SUM({S}{rc}8:{rc}10)/SUM({S}{uc}8:{uc}10),0)', fmt=AED)
    put(f'{col}{r + 4}', f'=IFERROR({col}{r + 2}/{col}{r + 3}-1,0)', fmt=PCT)
    put(f'{col}{r + 5}', f'=IFERROR({S}{uc}11/{S}{uc}7,0)', fmt=PCT)
    put(f'{col}{r + 6}', f'=SUM({S}{uc}8:{uc}10)', fmt=NUM)
    put(f'{col}{r + 7}', f'=N({S}{uc}11)', fmt=NUM)
ws.conditional_formatting.add(f'C{r + 4}:R{r + 4}', openpyxl.formatting.rule.CellIsRule(operator='greaterThan', formula=['0.2'], fill=fill_red))
put(f'B{r + 8}', ("Feb-25 order-level data shows 326.8 was already Maxima's full-price point (9 units); 285.3 was the "
                  "regular promoted price (225 units). Since Jan-26 (and mostly in Dec-25) every unit sells at 326.8, so the promo looks switched off. "
                  "Meanwhile Play 50 ml went down from ~272 to ~250. The premium widened from ~5% to ~30% and Maxima's "
                  "share of 50 ml units fell from ~16% to ~10%."), f_note, border=False, align=wrap)
ws.merge_cells(f'B{r + 8}:R{r + 8}')
ws.row_dimensions[r + 8].height = 42
maxima_end = r + 8

# ---------------------------------------------------------------- C. Ramadan
r = maxima_end + 2
section(r, 'C. Ramadan effect on lubes units (your remark: "sales drop in Ramadan")', 8)
header(r + 1, ['Year', 'Pre-Ramadan base', 'Base definition', 'Ramadan month(s)', 'Ramadan units',
               'Dip vs base', 'Post-Ramadan (Apr)', 'Apr vs base'])
put(f'B{r + 2}', '2025')
put(f'C{r + 2}', f'=AVERAGE({S}U17:V17)', fmt=NUM)
put(f'D{r + 2}', 'Jan-Feb avg')
put(f'E{r + 2}', 'Mar-25 (Ramadan ~1-30 Mar)')
put(f'F{r + 2}', f'={S}W17', fmt=NUM)
put(f'G{r + 2}', f'=F{r + 2}/C{r + 2}-1', fmt=PCT)
put(f'H{r + 2}', f'={S}X17', fmt=NUM)
put(f'I{r + 2}', f'=H{r + 2}/C{r + 2}-1', fmt=PCT)
put(f'B{r + 3}', '2026')
put(f'C{r + 3}', f'={S}AH17', fmt=NUM)
put(f'D{r + 3}', 'Jan')
put(f'E{r + 3}', 'Feb-Mar 26 avg (Ramadan ~18 Feb-19 Mar)')
put(f'F{r + 3}', f'=AVERAGE({S}AI17:AJ17)', fmt=NUM)
put(f'G{r + 3}', f'=F{r + 3}/C{r + 3}-1', fmt=PCT)
put(f'H{r + 3}', f'={S}AK17', fmt=NUM)
put(f'I{r + 3}', f'=H{r + 3}/C{r + 3}-1', fmt=PCT)
for rr in (r + 2, r + 3):
    ws[f'E{rr}'].alignment = wrap
    ws.row_dimensions[rr].height = 30
put(f'B{r + 4}', ("Confirmed: Ramadan takes 13-18% off lubes units, and April has not fully recovered to the Jan-26 base yet. "
                  "Ramadan 2027 is expected around 8 Feb - 9 Mar, so plan the pre-load and Eid visibility now."),
    f_note, border=False, align=wrap)
ws.merge_cells(f'B{r + 4}:I{r + 4}')
ws.row_dimensions[r + 4].height = 30
ram_end = r + 4

# ---------------------------------------------------------------- D. store reach
r = ram_end + 2
section(r, 'D. Store reach and velocity (Raw data: Location is only filled for Jan/Apr/May/Aug 2025)', 11)
header(r + 1, ['SKU', 'Stores Jan-25', 'Stores Apr-25', 'Stores May-25', 'Stores Aug-25',
               'Units/store Jan-25', 'Units/store Apr-25', 'Units/store May-25', 'Units/store Aug-25',
               'Reach Aug-25 (of 40 active)'])
reach = [  # stores selling >=1 unit (distinct Location), hardcoded from Raw data
    ('Play Cheeky Cherry', 8, [41, 35, 36, 39]),
    ('Play Feel', 9, [43, 37, 37, 39]),
    ('Play Tingle', 10, [39, 36, 36, 39]),
    ('Maxima Lubricant Gel', 11, [38, 35, 36, 37]),
    ('Play Massage 2-in-1 Aloe Vera', 13, [34, 34, 31, 36]),
    ('Play Massage 2-in-1 Ylang Ylang', 14, [14, 33, 31, 37]),
    ('Naturals Pure Intimate', 16, [0, 1, 0, 1]),
]
active = [62, 37, 37, 40]   # distinct stores with any Durex sale in that month
ucols = ['U', 'X', 'Y', 'AB']
row = r + 2
for name, s, st in reach:
    put(f'B{row}', name)
    for j, n in enumerate(st):
        put(f'{CL(3 + j)}{row}', n, f_input, fmt=NUM)
        put(f'{CL(7 + j)}{row}', f'=IFERROR({S}{ucols[j]}{s}/{CL(3 + j)}{row},0)', fmt=DEC)
    put(f'K{row}', f'=IFERROR(F{row}/F${row - (row - (r + 2)) + len(reach)},0)', fmt=PCT)
    row += 1
put(f'B{row}', 'Active Durex stores (any SKU)', f_bold, fill=fill_tot)
for j, n in enumerate(active):
    put(f'{CL(3 + j)}{row}', n, Font(name=FONT, size=11, bold=True, color='0000FF'), fmt=NUM, fill=fill_tot)
for col in 'GHIJK':
    put(f'{col}{row}', '', fill=fill_tot)
# fix reach formula denominators to point at the active row
for rr in range(r + 2, row):
    ws[f'K{rr}'] = f'=IFERROR(F{rr}/F${row},0)'
put(f'B{row + 1}', ("Blue = hardcoded from 'Raw data' (distinct Location with >=1 unit sold). Jan-25 includes 26 "
                    "'Talabat Mart - Wellness' stores that sold only 73 Durex units in total and drop out after January. "
                    "In 2025 the core SKUs sold in 90-100% of stores, so ranging was not the issue then. "
                    "There is no store-level data for 2026, so ask Talabat for 2026 in-stock % by dark store before calling "
                    "the decline a demand problem."), f_note, border=False, align=wrap)
ws.merge_cells(f'B{row + 1}:K{row + 1}')
ws.row_dimensions[row + 1].height = 55
reach_end = row + 1

# ---------------------------------------------------------------- E. remarks validation
r = reach_end + 2
section(r, 'E. Your remarks, checked against the data', 6)
header(r + 1, ['Your remark (col A, Lubes SKU level summary)', 'What the data says', 'Verdict', 'So what / action'])
ws.merge_cells(start_row=r + 1, start_column=6, end_row=r + 1, end_column=7)
remarks = [
    ('Naturals 100 ml: consistent in-stock / sales issue, fixed price, never changes, still not in stock',
     'Sold in 9 of 16 months, 17 units total, in at most 1 store a month. ASP is flat at 667.9 (no promo ever). '
     'At ~6.7 per ml it is the most expensive lube per ml (50 ml ~5.0, 200 ml ~2.9).',
     'Confirmed',
     'Either get it properly ranged and in stock across all ~40 stores or swap the slot for a SKU that sells '
     '(e.g. Extra Sensitive 100 ml from Amazon). Do not let a dead listing hold the 100 ml slot.'),
    ('Sales drop in Ramadan',
     'Units fell 18% in Mar-25 vs Jan-Feb. In 2026 Feb-Mar averaged 13% below Jan.',
     'Confirmed',
     'Part of the 2026 "decline" is seasonal. Compare April or ex-Ramadan months when sizing the real gap.'),
    ('200 ml has 3 variants (Aloe, Ylang Ylang, Guarana). Guarana is not on Talabat',
     'The two listed 200 ml SKUs are 20% of lubes revenue at ~2x the ASP of 50 ml, with ~3 units/store/month.',
     'Range gap',
     'List Guarana 200 ml. See sizing in F.'),
    ('Play has 5 variants in 50 ml; Heat and Strawberry are not on Talabat. 8 Play variants in total, 5 on Talabat',
     'The 3 listed Play 50 ml SKUs are ~70% of lubes revenue and growing +18% to +56% in units YoY. '
     'Maxima 50 ml is not a Play SKU, so the 5 of 8 count holds.',
     'Range gap',
     'List Heat and Strawberry 50 ml. The category is still growing YoY, so new variants have room. See F.'),
    ('Aloe Vera sells more than Ylang Ylang in 200 ml',
     "Aloe 497 vs Ylang 399 units JA'26 (+25%), and 1,957 vs 1,505 lifetime. But Aloe is ~16% below its H2-25 "
     "level in April while Ylang is ~11% below.",
     'Confirmed, with a watch-out',
     'Aloe is the bigger SKU but also the bigger leak right now. Check its 2026 in-stock first.'),
    ('Big price gap vs Amazon: Play on Talabat is ~1.78x Amazon today',
     'Cannot be checked from this file (Talabat only). On Talabat the Play 50 ml ASP is ~8% lower YoY (272 to 250), '
     'so any widening gap came from Amazon going down, not Talabat going up. Maxima is the exception: +14% on Talabat.',
     'Plausible, needs evidence',
     'Pull Amazon price history (Keepa) for the same SKUs to show when the gap opened, then line it up with the '
     'Talabat units dip. Use it as the price-parity ask in the JBP.'),
    ('Extra Sensitive 100 ml is on Amazon, not on Talabat. Amazon has many more variants not live on Talabat',
     'Talabat has 7 lube SKUs listed, of which 6 actually sell. The 100 ml slot is effectively empty.',
     'Range gap',
     'Build the must-have list from Amazon best-sellers and pitch it to Talabat as a range extension.'),
]
row = r + 2
for rem, data, verdict, act in remarks:
    put(f'B{row}', rem, align=wrap)
    put(f'C{row}', data, align=wrap)
    put(f'D{row}', verdict, f_bold, align=center,
        fill=fill_grn if verdict.startswith('Confirmed') else (fill_red if verdict == 'Range gap' else fill_in))
    put(f'E{row}', act, align=wrap)
    ws.merge_cells(f'E{row}:G{row}')
    ws.row_dimensions[row].height = 75
    row += 1
rem_end = row - 1

# ---------------------------------------------------------------- F. range-gap sizing
r = rem_end + 2
section(r, 'F. Range-gap sizing (indicative): edit the yellow inputs', 8)
header(r + 1, ['Missing SKU', 'Benchmark SKU', 'Benchmark units / mo (JA\'26)', 'Velocity vs benchmark',
               'Cannibalisation', 'ASP', 'Incremental rev / mo', 'Incremental rev / yr'])
gaps = [
    ('Play Heat 50 ml', 'Play Tingle (weakest Play 50 ml)', 10, 0.8, 0.4),
    ('Play Strawberry 50 ml', 'Play Tingle (weakest Play 50 ml)', 10, 0.8, 0.4),
    ('Play Guarana 200 ml', 'Play Ylang Ylang 200 ml', 14, 0.8, 0.3),
]
row = r + 2
g_first = row
for name, bench, s, vel, can in gaps:
    put(f'B{row}', name)
    put(f'C{row}', bench)
    put(f'D{row}', f'=SUM({S}AH{s}:AK{s})/4', fmt=NUM)
    put(f'E{row}', vel, f_input, fmt='0%', fill=fill_in)
    put(f'F{row}', can, f_input, fmt='0%', fill=fill_in)
    put(f'G{row}', f'=IFERROR(SUM({S}P{s}:S{s})/SUM({S}AH{s}:AK{s}),0)', fmt=AED)
    put(f'H{row}', f'=D{row}*E{row}*(1-F{row})*G{row}', fmt=AED)
    put(f'I{row}', f'=H{row}*12', fmt=AED)
    row += 1
put(f'B{row}', 'Total', f_bold, fill=fill_tot)
for col in 'CDEFG':
    put(f'{col}{row}', '', fill=fill_tot)
put(f'H{row}', f'=SUM(H{g_first}:H{row - 1})', f_bold, fmt=AED, fill=fill_tot)
put(f'I{row}', f'=SUM(I{g_first}:I{row - 1})', f_bold, fmt=AED, fill=fill_tot)
put(f'J{row}', f"=IFERROR(I{row}/SUM({S}I17:N17,{S}P17:S17)*10/12,0)", fmt=PCT, border=False)
put(f'K{row}', '<- as % of the current lubes annual run rate', f_note, border=False)
put(f'B{row + 1}', ("Assumptions are illustrative and set by me, not taken from the data. Velocity: a new variant sells at 80% of "
                    "the benchmark SKU. Cannibalisation: the share of its units taken from existing Play SKUs. The run-rate "
                    "denominator is the last 10 months (Jul-25 to Apr-26), annualised."), f_note, border=False, align=wrap)
ws.merge_cells(f'B{row + 1}:I{row + 1}')
ws.row_dimensions[row + 1].height = 42

# ---------------------------------------------------------------- layout
widths = {'A': 2, 'B': 34, 'C': 30, 'D': 16, 'E': 34, 'F': 13, 'G': 13, 'H': 13, 'I': 13, 'J': 12, 'K': 12,
          'L': 13, 'M': 13, 'N': 12, 'O': 13, 'P': 12, 'Q': 55, 'R': 11}
for k, v in widths.items():
    ws.column_dimensions[k].width = v
ws.sheet_view.showGridLines = False
ws.freeze_panes = 'C4'

wb.save(out)
print('saved', out)
