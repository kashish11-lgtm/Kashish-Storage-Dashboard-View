import sys, openpyxl
from openpyxl.styles import Font
src = openpyxl.load_workbook(sys.argv[1], data_only=True)['Lubes SKU level summary']
wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Lubes SKU level summary'
for row in src.iter_rows():
    for c in row:
        v = c.value
        if isinstance(v, str) and v.startswith('#'):
            v = None
        if v is not None:
            n = ws.cell(row=c.row, column=c.column, value=v)
            n.font = Font(name='Aptos Narrow', size=11, bold=c.font.bold)
            n.number_format = c.number_format
for k, d in src.column_dimensions.items():
    ws.column_dimensions[k].width = d.width
ws['A2'] = 'Values copy of the pivot in Talabat_Durex_cleaned_Analysis.xlsx (same cell positions), so the Deep Dive formulas resolve.'
ws['A2'].font = Font(name='Aptos Narrow', size=9, italic=True)
wb.save(sys.argv[2])
